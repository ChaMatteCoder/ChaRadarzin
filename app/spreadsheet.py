from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.models import Catalog, Product, ProductLink


PRODUCT_COLUMNS = (
    "produto_id",
    "ativo",
    "produto",
    "modelo_exato",
    "preco_alvo",
    "pagamento",
)
LINK_COLUMNS = (
    "produto_id",
    "loja",
    "url",
    "vendedor_esperado",
    "variante",
)


class SpreadsheetValidationError(ValueError):
    def __init__(self, issues: list[str]) -> None:
        self.issues = issues
        super().__init__("Planilha invalida:\n- " + "\n- ".join(issues))


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _headers(sheet: Any, required: tuple[str, ...]) -> tuple[int, dict[str, int]]:
    required_set = set(required)
    for row in range(1, min(sheet.max_row, 20) + 1):
        headers = {
            _text(cell.value).lower(): index
            for index, cell in enumerate(sheet[row], start=1)
            if _text(cell.value)
        }
        if required_set.issubset(headers):
            return row, headers
    return 1, {
        _text(cell.value).lower(): index
        for index, cell in enumerate(sheet[1], start=1)
        if _text(cell.value)
    }


def _require_columns(
    sheet: Any, required: tuple[str, ...], issues: list[str]
) -> tuple[int, dict[str, int]]:
    header_row, headers = _headers(sheet, required)
    missing = [column for column in required if column not in headers]
    if missing:
        issues.append(f"Aba {sheet.title}: colunas ausentes: {', '.join(missing)}")
    return header_row, headers


def _cell(sheet: Any, row: int, headers: dict[str, int], column: str) -> Any:
    index = headers.get(column)
    return None if index is None else sheet.cell(row=row, column=index).value


def _parse_active(value: Any, label: str, issues: list[str]) -> bool:
    normalized = _text(value).casefold()
    if normalized in {"sim", "s", "true", "1", "yes"}:
        return True
    if normalized in {"nao", "não", "n", "false", "0", "no"}:
        return False
    issues.append(f"{label}: 'ativo' deve ser Sim ou Nao")
    return False


def _parse_price(value: Any, label: str, issues: list[str]) -> Decimal:
    if value is None or _text(value) == "":
        issues.append(f"{label}: 'preco_alvo' e obrigatorio")
        return Decimal("0")
    try:
        if isinstance(value, str):
            normalized = value.replace("R$", "").replace(" ", "")
            if "," in normalized:
                normalized = normalized.replace(".", "").replace(",", ".")
            return Decimal(normalized)
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        issues.append(f"{label}: 'preco_alvo' deve ser numerico")
        return Decimal("0")


def load_catalog(path: str | Path) -> Catalog:
    workbook_path = Path(path)
    if not workbook_path.exists():
        raise SpreadsheetValidationError([f"Arquivo nao encontrado: {workbook_path}"])

    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    try:
        issues: list[str] = []
        missing_sheets = [name for name in ("Produtos", "Links") if name not in workbook.sheetnames]
        if missing_sheets:
            raise SpreadsheetValidationError(
                [f"Abas obrigatorias ausentes: {', '.join(missing_sheets)}"]
            )

        products_sheet = workbook["Produtos"]
        links_sheet = workbook["Links"]
        product_header_row, product_headers = _require_columns(
            products_sheet, PRODUCT_COLUMNS, issues
        )
        link_header_row, link_headers = _require_columns(links_sheet, LINK_COLUMNS, issues)
        if issues:
            raise SpreadsheetValidationError(issues)

        products: list[Product] = []
        product_ids: set[str] = set()
        for row in range(product_header_row + 1, products_sheet.max_row + 1):
            if all(_cell(products_sheet, row, product_headers, column) is None for column in PRODUCT_COLUMNS):
                continue
            label = f"Aba Produtos, linha {row}"
            product_id = _text(_cell(products_sheet, row, product_headers, "produto_id")).upper()
            name = _text(_cell(products_sheet, row, product_headers, "produto"))
            exact_model = _text(_cell(products_sheet, row, product_headers, "modelo_exato"))
            payment = _text(_cell(products_sheet, row, product_headers, "pagamento")).upper()
            if not product_id:
                issues.append(f"{label}: 'produto_id' e obrigatorio")
            elif product_id in product_ids:
                issues.append(f"{label}: produto_id duplicado: {product_id}")
            else:
                product_ids.add(product_id)
            if not name:
                issues.append(f"{label}: 'produto' e obrigatorio")
            if not exact_model:
                issues.append(f"{label}: 'modelo_exato' e obrigatorio")
            if not payment:
                issues.append(f"{label}: 'pagamento' e obrigatorio")
            products.append(
                Product(
                    product_id=product_id,
                    active=_parse_active(
                        _cell(products_sheet, row, product_headers, "ativo"), label, issues
                    ),
                    name=name,
                    exact_model=exact_model,
                    target_price=_parse_price(
                        _cell(products_sheet, row, product_headers, "preco_alvo"), label, issues
                    ),
                    payment_method=payment,
                )
            )

        if not products:
            issues.append("Aba Produtos: informe ao menos um produto")

        links: list[ProductLink] = []
        seen_links: set[tuple[str, str, str]] = set()
        for row in range(link_header_row + 1, links_sheet.max_row + 1):
            if all(_cell(links_sheet, row, link_headers, column) is None for column in LINK_COLUMNS):
                continue
            label = f"Aba Links, linha {row}"
            product_id = _text(_cell(links_sheet, row, link_headers, "produto_id")).upper()
            store = _text(_cell(links_sheet, row, link_headers, "loja"))
            url = _text(_cell(links_sheet, row, link_headers, "url"))
            seller = _text(_cell(links_sheet, row, link_headers, "vendedor_esperado"))
            variant = _text(_cell(links_sheet, row, link_headers, "variante"))
            for field, value in (
                ("produto_id", product_id),
                ("loja", store),
                ("url", url),
                ("vendedor_esperado", seller),
                ("variante", variant),
            ):
                if not value:
                    issues.append(f"{label}: '{field}' e obrigatorio")
            if product_id and product_id not in product_ids:
                issues.append(f"{label}: produto_id inexistente em Produtos: {product_id}")
            key = (product_id, store.casefold(), url)
            if all(key) and key in seen_links:
                issues.append(f"{label}: link duplicado para {product_id} / {store}")
            seen_links.add(key)
            links.append(
                ProductLink(
                    product_id=product_id,
                    store=store,
                    url=url,
                    expected_seller=seller,
                    variant=variant,
                )
            )

        active_ids = {product.product_id for product in products if product.active}
        linked_ids = {link.product_id for link in links}
        for product_id in sorted(active_ids - linked_ids):
            issues.append(f"Produto ativo {product_id} nao possui nenhum link cadastrado")

        if issues:
            raise SpreadsheetValidationError(issues)
        return Catalog(products=tuple(products), links=tuple(links))
    finally:
        workbook.close()
